from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from ai_agent_template.api.app import create_app
from ai_agent_template.api.routes import commerce as routes
from ai_agent_template.commerce.demo import manual, seed_structured
from ai_agent_template.commerce.listing import ListingContent, rule_issues
from ai_agent_template.commerce.orders import after_sales
from ai_agent_template.commerce.repository import CommerceRepository
from ai_agent_template.commerce.retrieval import CommerceScope, scope
from ai_agent_template.commerce.schemas import (
    AfterSalesRequest,
    ChatRequest,
    ListingRequest,
    Principal,
    Product,
    ProductVariant,
)
from ai_agent_template.commerce.service import CommerceService
from ai_agent_template.core.settings import Settings
from ai_agent_template.rag.evidence import GroundedClaim


@pytest.fixture
def system(tmp_path, monkeypatch):
    settings = Settings(
        _env_file=None,
        commerce_database_url=f"sqlite:///{tmp_path / 'business.sqlite3'}",
        rag_data_path=str(tmp_path / "rag"),
        commerce_llm_mode="offline",
        commerce_mode="demo",
        commerce_principals={},
        qdrant_url=None,
        rate_limit_enabled=False,
        auth_api_keys=[],
        openai_api_key=None,
        groq_api_key=None,
    )
    service = CommerceService(settings)
    from ai_agent_template.commerce.offline import FakeCommercePolicy, FakeCommerceWriter

    service._model, service._policy = FakeCommerceWriter(), FakeCommercePolicy(service.repository)
    monkeypatch.setattr(routes, "get_settings", lambda: settings)
    app = create_app()
    app.dependency_overrides[routes.get_commerce_service] = lambda: service
    client = TestClient(app)
    client.headers["Authorization"] = "Bearer demo-operator"
    yield service, client, settings
    service.close()


def principal(store="demo-store", customer=False):
    return Principal(
        store_id=store,
        subject="tester",
        role="customer" if customer else "operator",
        customer_id="demo-customer" if customer else None,
    )


def test_catalog_twenty_variants_and_exact_attribute_filter(system):
    service, client, settings = system
    response = client.get("/v1/commerce/products")
    assert response.status_code == 200
    assert len(response.json()["variants"]) == 20
    sku = client.get("/v1/commerce/products", params={"sku": "DEMO-SERUM-30"}).json()["variants"]
    assert len(sku) == 1 and sku[0]["attributes"]["capacity"] == "30 ml"
    result = client.get(
        "/v1/commerce/products", params={"q": "精华", "attribute": "capacity", "value": "60 ml"}
    ).json()
    assert [item["sku"] for item in result["variants"]] == ["DEMO-SERUM-60"]
    assert service.repository.health()
    settings.commerce_principals["only-us"] = dict(
        store_id="demo-store", subject="restricted-operator", role="operator", marketplaces=["US"]
    )
    client.headers["Authorization"] = "Bearer only-us"
    assert client.get("/v1/commerce/products", params={"market": "DE"}).status_code == 403
    assert (
        client.get("/v1/commerce/search", params={"q": "product", "market": "DE"}).status_code
        == 403
    )
    assert (
        client.post("/v1/commerce/chat", json={"message": "product", "market": "DE"}).status_code
        == 403
    )


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_spreadsheet_structure_not_embedded(system, extension):
    service, client, _ = system
    headers = ["sku", "product_id", "name_en", "capacity", "price", "stock", "description"]
    row = [
        "IMPORT-1",
        "import-product",
        "Imported product",
        "75 ml",
        "987.65",
        "4321",
        "Use on dry surfaces. No certification supplied.",
    ]
    if extension == "csv":
        data = (",".join(headers) + "\n" + ",".join(row)).encode()
    else:
        book = Workbook()
        book.active.append(headers)
        book.active.append(row)
        buffer = BytesIO()
        book.save(buffer)
        data = buffer.getvalue()
    response = client.post(
        "/v1/commerce/catalog/import", files={"file": ("catalog." + extension, data)}
    )
    assert response.status_code == 200, response.text
    variant = service.repository.get(ProductVariant, "demo-store", "IMPORT-1")
    assert str(variant.price) == "987.65" and variant.stock == 4321
    doc = response.json()["documents"][0]["document"]
    records = [
        service.rag("demo-store").repository.resolve(item)[0] for item in doc["evidence_ids"]
    ]
    assert all("987.65" not in item.text and "4321" not in item.text for item in records)


async def upload(
    service,
    store="demo-store",
    sku="DEMO-SERUM-30",
    text="SKU DEMO-SERUM-30. Capacity 30 ml. Fragrance free.",
    document_id=None,
    **kwargs,
):
    return await service.upload(
        principal(store), sku, text.encode(), "manual.txt", "US", "en", document_id, **kwargs
    )


async def test_sku_scope_and_cross_store_isolation_before_retrieval(system):
    service, _, _ = system
    seed_structured(service.repository, "another-store")
    first = await upload(service)
    await upload(service, "another-store", text="SKU DEMO-SERUM-30. Foreign confidential note.")
    token = scope.set(CommerceScope(principal(), "US", "en", ["DEMO-SERUM-30"]))
    try:
        sources = service.retriever.search("Foreign confidential note", mode="hybrid")
        assert sources
        assert all(item.metadata["store_id"] == "demo-store" for item in sources)
        assert all("Foreign confidential" not in item.content for item in sources)
        assert any(
            item.evidence.document_id == first["document"]["document_id"]
            for item in sources
            if item.evidence
        )
    finally:
        scope.reset(token)


async def test_versions_duplicate_expiry_and_original(system):
    service, client, _ = system
    original = await upload(service, text="SKU DEMO-SERUM-30. Old ingredient is Alpha.")
    doc = original["document"]
    duplicate = await upload(service, text="SKU DEMO-SERUM-30. Old ingredient is Alpha.")
    assert duplicate["index"]["stats"]["embedding_cache_misses"] == 0
    assert duplicate["document"]["document_version"] == 1
    updated = await upload(
        service, text="SKU DEMO-SERUM-30. New ingredient is Beta.", document_id=doc["document_id"]
    )
    assert updated["document"]["document_version"] == 2
    assert (
        client.get("/v1/commerce/evidence/" + doc["evidence_ids"][0]).json()["status"]
        == "historical"
    )
    assert client.get(f"/v1/commerce/documents/{doc['document_id']}/1/file").content.endswith(
        b"Alpha."
    )
    token = scope.set(CommerceScope(principal(), "US", "en", ["DEMO-SERUM-30"]))
    try:
        assert all("Alpha" not in item.content for item in service.retriever.search("ingredient"))
    finally:
        scope.reset(token)
    expired = await upload(
        service,
        text="Expired certification XYZ.",
        document_id=doc["document_id"],
        expires_at="2020-01-01T00:00:00+00:00",
    )
    token = scope.set(CommerceScope(principal(), "US", "en", ["DEMO-SERUM-30"]))
    try:
        assert all("XYZ" not in item.content for item in service.retriever.search("XYZ"))
    finally:
        scope.reset(token)
    assert (
        client.get("/v1/commerce/evidence/" + expired["document"]["evidence_ids"][0]).json()[
            "status"
        ]
        == "expired"
    )


@pytest.mark.parametrize(
    "language,question",
    [
        ("zh", "DEMO-SERUM-30 的容量是多少？"),
        ("en", "What is the capacity of DEMO-SERUM-30?"),
        ("es", "¿Cuál es la capacidad de DEMO-SERUM-30?"),
    ],
)
async def test_multilingual_chat_and_citations(system, language, question):
    service, _, _ = system
    await upload(service)
    result = await service.chat(principal(), ChatRequest(message=question, language=language))
    assert result["mode"] == "offline_demo"
    assert result["citations"] and result["verification"]["supported"]
    assert "30" in result["answer"]
    assert result["query_plan"]["intent"] == "product_faq"


def test_listing_rules_reject_unbound_efficacy():
    content = ListingContent(
        title="FDA-approved cure",
        bullets=["Invented"],
        description="cures disease",
        claims=[GroundedClaim(text="Some different text", evidence_ids=["E1"])],
    )
    issues = rule_issues(content, ListingRequest(sku="A", market="US"))
    assert any("未绑定" in item for item in issues)
    assert any("敏感" in item for item in issues)
    assert any(
        "未配置" in item for item in rule_issues(content, ListingRequest(sku="A", market="ES"))
    )


async def test_listing_draft_checked_and_pending_review(system):
    service, client, _ = system
    await upload(service)
    response = client.post(
        "/v1/commerce/listings",
        json={"sku": "DEMO-SERUM-30", "language": "en", "market": "US", "platform": "amazon"},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["review_status"] == "pending_review", result
    assert result["generation_mode"] == "offline_extractive_demo"
    assert result["facts"]
    client.headers["Authorization"] = "Bearer demo-customer"
    assert client.post("/v1/commerce/listings", json={"sku": "DEMO-SERUM-30"}).status_code == 403


def test_order_authorization_and_no_anonymous_access(system):
    _, client, _ = system
    client.headers.pop("Authorization")
    assert client.get("/v1/commerce/orders/DEMO-ORDER-1").status_code == 401
    client.headers["Authorization"] = "Bearer demo-customer"
    response = client.get("/v1/commerce/orders/DEMO-ORDER-1")
    assert response.status_code == 200
    assert response.json()["order"]["status"] == "completed"
    assert response.json()["actions_executed"] == []
    assert client.get("/v1/commerce/orders/DEMO-ORDER-3").status_code == 404
    assert (
        client.get(
            "/v1/commerce/orders/DEMO-ORDER-1",
            headers={"x-store-id": "another-store", "x-customer-id": "other-customer"},
        ).json()["order"]["customer_id"]
        == "demo-customer"
    )


def test_policy_eligibility_no_automatic_refund_and_untrusted_date(system):
    service, _, _ = system
    accepted = after_sales(
        service.repository,
        principal(customer=True),
        "DEMO-ORDER-1",
        AfterSalesRequest(reason="damaged"),
    )
    assert accepted["eligible_for_review"] is True and accepted["actions_executed"] == []
    assert accepted["case"]["status"] == "pending_human_review"
    denied = after_sales(
        service.repository,
        principal(),
        "DEMO-ORDER-1",
        AfterSalesRequest(reason="changed_mind", opened=True),
    )
    assert denied["eligible_for_review"] is False
    unknown = after_sales(
        service.repository, principal(), "DEMO-ORDER-1", AfterSalesRequest(reason="changed_mind")
    )
    assert unknown["eligible_for_review"] is None
    tampered = after_sales(
        service.repository,
        principal(),
        "DEMO-ORDER-1",
        AfterSalesRequest(reason="damaged", received_at="2030-01-01T00:00:00+00:00"),
    )
    assert tampered["eligible_for_review"] is None


def test_production_requires_postgres_and_disables_demo_identity(system, monkeypatch):
    _, client, settings = system
    with pytest.raises(ValueError, match="PostgreSQL"):
        CommerceRepository("sqlite:///:memory:", production=True)
    monkeypatch.setattr(
        routes, "get_settings", lambda: settings.model_copy(update={"commerce_mode": "production"})
    )
    assert client.get("/v1/commerce/orders/DEMO-ORDER-1").status_code == 401


def test_missing_live_key_is_explicit(system):
    service, _, settings = system
    service.settings = settings.model_copy(
        update={"commerce_llm_mode": "live", "model_provider": "openai", "openai_api_key": None}
    )
    with pytest.raises(ValueError, match="真实 LLM 未配置"):
        service.models()


def test_sse_and_invalid_import(system):
    _, client, _ = system
    stream = client.post(
        "/v1/commerce/chat/stream",
        json={"message": "capacity DEMO-SERUM-30", "skus": ["DEMO-SERUM-30"]},
    )
    assert stream.status_code == 200, stream.text
    assert "event: done" in stream.text and "citations" in stream.text
    response = client.post(
        "/v1/commerce/catalog/import", files={"file": ("bad.csv", b"sku,name_en\nA,Name")}
    )
    assert response.status_code == 400


def test_document_parser_docx_and_image_disabled(system):
    _, client, _ = system
    from docx import Document

    word = Document()
    word.add_paragraph("SKU DEMO-SERUM-30 capacity 30 ml.")
    data = BytesIO()
    word.save(data)
    response = client.post(
        "/v1/commerce/documents",
        data={"sku": "DEMO-SERUM-30"},
        files={"file": ("manual.docx", data.getvalue())},
    )
    assert response.status_code == 200, response.text
    response = client.post(
        "/v1/commerce/documents",
        data={"sku": "DEMO-SERUM-30"},
        files={"file": ("label.png", b"invalid")},
    )
    assert response.status_code == 400


async def test_bilingual_manual_claim_binding_and_multiturn_variant(system):
    service, _, _ = system
    variant = service.repository.get(ProductVariant, "demo-store", "DEMO-SERUM-60")
    product = service.repository.get(Product, "demo-store", variant.product_id)
    await upload(service, sku=variant.sku, text=manual(variant, product))
    result = await service.chat(
        principal(),
        ChatRequest(message="DEMO-SERUM-60 有香精吗？", language="zh", thread_id="variants"),
    )
    assert result["verification"]["supported"], result["verification"]
    follow_up = await service.chat(
        principal(), ChatRequest(message="那小容量版本呢？", language="zh", thread_id="variants")
    )
    assert follow_up["query_plan"]["skus"] == ["DEMO-SERUM-30"]
    assert all(
        item["evidence"]["metadata"].get("sku") == "DEMO-SERUM-30"
        for item in follow_up["citations"]
    )


def test_first_concurrent_service_initialization_is_single_instance(system, monkeypatch):
    service, _, settings = system
    from concurrent.futures import ThreadPoolExecutor

    calls = []
    routes._build_commerce_service.cache_clear()

    def build(current_settings):
        calls.append(current_settings)
        return service

    monkeypatch.setattr(routes, "CommerceService", build)
    with ThreadPoolExecutor(max_workers=8) as pool:
        values = list(pool.map(lambda _: routes.get_commerce_service(), range(16)))
    assert len(calls) == 1 and all(value is service for value in values)
    routes._build_commerce_service.cache_clear()


def test_plain_chinese_sentence_keeps_trailing_citation():
    from ai_agent_template.rag.evidence import parse_claims

    claims = parse_claims("成分是水和甘油。 [E1]\nNo certification supplied. [E2]")
    assert [item.evidence_ids for item in claims] == [["E1"], ["E2"]]


def test_production_environment_never_uses_default_demo_token(system, monkeypatch):
    _, client, settings = system
    monkeypatch.setattr(
        routes,
        "get_settings",
        lambda: settings.model_copy(update={"environment": "production", "commerce_mode": "demo"}),
    )
    assert client.get("/v1/commerce/orders/DEMO-ORDER-1").status_code == 401


async def test_commerce_authority_gate_rejects_document_price(system):
    from ai_agent_template.commerce.verification import CommerceVerifier
    from ai_agent_template.rag.evidence import ClaimAssessment, VerificationReport
    from ai_agent_template.rag.schemas import EvidenceRecord

    service, _, _ = system

    class Judge:
        async def assess(self, claims, records):
            return VerificationReport(
                level="semantic",
                supported=True,
                checks=[
                    ClaimAssessment(
                        claim_index=0,
                        claim=claims[0].text,
                        evidence_ids=["E1"],
                        verdict="supported",
                        reason="copied number",
                    )
                ],
            )

    record = EvidenceRecord(
        evidence_id="x",
        document_id="x",
        document_version=1,
        chunk_id="x",
        source_filename="old-price.pdf",
        text="DEMO-SERUM-30 price is 99 USD.",
        text_hash="x",
        metadata={"store_id": "demo-store", "sku": "DEMO-SERUM-30", "market": "US"},
    )
    token = scope.set(CommerceScope(principal(), "US", "en", ["DEMO-SERUM-30"]))
    try:
        result = await CommerceVerifier(service.repository, Judge()).assess(
            [GroundedClaim(text="DEMO-SERUM-30 price is 99 USD.", evidence_ids=["E1"])],
            {"E1": record},
        )
        assert not result.supported and "authoritative" in result.checks[0].reason
    finally:
        scope.reset(token)
